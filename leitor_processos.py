#!/usr/bin/env python3
"""
Leitor de notificações trabalhistas (PDF escaneado) -> planilha de audiências.

A planilha gerada segue o modelo "Audiências 2026":
STATUS | DATA | HORA | HOSPITAL | RECLAMANTE | PROCESSO | VARA | MODALIDADE |
AUDIÊNCIAS | TESTEMUNHAS | OBSERVAÇÃO | LINK E SENHA
O que não for encontrado no documento fica em branco.

Regras (baseadas nos PDFs de exemplo do TRT-1):
  * Cada notificação tem uma página principal (nº do processo, RECLAMANTE,
    DESTINATÁRIO, data/hora, tipo da audiência) e, em geral, uma página de envelope.
  * HOSPITAL = quem foi notificado (campo DESTINATÁRIO). Uma linha por PROCESSO:
    se o mesmo processo chegar para várias empresas da rede (mesmo em PDFs
    diferentes), os nomes ficam juntos na mesma célula, um por linha.
  * Mesmo reclamante com processos diferentes = linhas diferentes (só é avisado).

Também é usado pela versão web (servidor.py), que roda dentro do ERP.

Uso:
    python leitor_processos.py pasta_com_pdfs                 # nome do Excel sai pelas datas
    python leitor_processos.py pasta_com_pdfs -o semana.xlsx --salvar-texto
    python leitor_processos.py pasta_com_pdfs --tesseract-cmd "C:/Program Files/Tesseract-OCR/tesseract.exe"

Dependências: pip install pymupdf pytesseract pillow openpyxl rapidfuzz (rapidfuzz é opcional)
Tesseract instalado no computador, com o idioma português (por).
"""
import argparse
import re
import sys
import unicodedata
from collections import Counter
from datetime import datetime, time
from pathlib import Path

# --------------------------------------------------------------------------
# Configuração
# --------------------------------------------------------------------------
STATUS_PADRAO = "Agendada"  # coluna STATUS. Deixe "" para sair em branco.

# Nome que vai na coluna HOSPITAL. A regra olha se o TRECHO aparece no nome do
# destinatário (sem acento, maiúsculo); vale a primeira que bater. Se nenhuma
# bater, usa o nome do documento sem "LTDA"/"S.A.". Complete conforme precisar.
APELIDOS = [
    ("CASA DE PORTUGAL- HOSPITAL GERAL", "HOSPITAL CASA DE PORTUGAL"),
    ("ARRENDATARIA", "HOSPITAL CASA ARRENDATARIA"),
    ("HOSPITAL DO CANCER", "HOSPITAL CASA HOSPITAL DO CANCER"),
    ("PRONTOCOR", "PRONTO SOCORRO CLINICO PRONTOCOR"),
    ("ILHA DO GOVERNADOR", "HOSPITAL ILHA DO GOVERNADOR"),
    ("3 D DIAGNOSE", "3D DIAGNOSE"),
    ("3D DIAGNOSE", "3D DIAGNOSE"),
]

# --------------------------------------------------------------------------
# Padrões
# --------------------------------------------------------------------------
# Nº CNJ. Tolera espaço que o OCR às vezes coloca no meio ("0101 308-58...") e
# ':' ou ',' lidos no lugar do ponto ("...2026.5:01.0012").
_SEP = r"\s?[.:,]\s?"
CNJ_RE = re.compile(
    rf"(?<!\d)(\d{{3,4}}\s?\d{{3,4}})\s?-\s?(\d{{2}}){_SEP}(\d{{4}}){_SEP}(\d){_SEP}(\d{{2}}){_SEP}(\d{{4}})(?!\d)"
)
DATA_HORA_RE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})\s*(?:AS\s*)?(\d{1,2})\s*[:H]\s*(\d{2})")
DATA_RE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
RECLAMANTE_RE = re.compile(r"^\W*RECLAMANTE\s*:\s*(.*)$")
RECLAMADO_RE = re.compile(r"^\W*RECLAMADO\s*:\s*(.*)$")
DEST_RE = re.compile(r"DESTINAT\w{0,6}\s*(?:\(\s*S\s*\))?\s*:\s*(.*)$")
OUTROS_RE = re.compile(r"\bE\s+OUTROS\s*\(?\s*(\d+)\s*\)?")
SUFIXO_RE = re.compile(r"\b(LTDA|S/?A|S\.A\.?|EIRELI|EPP|ME)\W*$")
ENDERECO_RE = re.compile(r"^\W*(RUA|AV|AVENIDA|TRAVESSA|PRACA|ESTRADA|LARGO|ALAMEDA)\b")
TOKEN_OK = re.compile(r"^[A-Z0-9(][A-Z0-9\-/.,&'()]*$")
SINAIS_ENVELOPE = ("POSTAG", "DEVOLU", "CORREIOS", "E-CARTA", "CDIP")
VARA_RE = re.compile(r"(\d{1,3})\s*[AO°*.]{0,2}\s+VARA\s+DO\s+TRABALHO")
ENDERECO_FORUM_RE = re.compile(
    r"\b(RUA|AVENIDA|AV)\.?\s+((?:DO|DA|DE|DOS|DAS)\s+)?([A-Z]+(?:\s+[A-Z]+){0,2}),\s*"
    r"(?:N[O°.]*\s*)?(\d{1,4})\s*,?\s*(?:(\d{1,2})\s*[AO°*]?\s*ANDAR)?"
)
ZOOM_RE = re.compile(
    r"https?://[A-Za-z0-9\-]+\.zoom\.us/(?:j|my)/[A-Za-z0-9._\-]+(?:\?pwd=[A-Za-z0-9]+)?", re.I
)
ID_REUNIAO_RE = re.compile(r"ID\s+DA\s+REUNIAO\s*(?:\(\s*PMI\s*\))?\s*[:.]?\s*((?:\d[\s.]?){9,12})", re.I)
SENHA_RE = re.compile(r"SENHA(?:\s+DE\s+ACESSO)?\s*[:.]?\s*([A-Za-z0-9]{4,12})", re.I)
PARTICULAS = {"DE", "DA", "DO", "DAS", "DOS", "E"}


# --------------------------------------------------------------------------
# Utilidades
# --------------------------------------------------------------------------
def sem_acentos(txt: str) -> str:
    txt = unicodedata.normalize("NFKD", txt)
    return "".join(c for c in txt if not unicodedata.combining(c))


def normalizar(txt: str) -> str:
    return re.sub(r"\s+", " ", sem_acentos(txt)).strip().upper()


def linhas_normalizadas(texto: str) -> list:
    return [re.sub(r"[ \t]+", " ", l).strip() for l in sem_acentos(texto).upper().splitlines()]


def razao_maiuscula(s: str) -> float:
    letras = [c for c in s if c.isalpha()]
    return sum(c.isupper() for c in letras) / len(letras) if letras else 0.0


def limpar_ruido(txt: str) -> str:
    """Tira 'lixo' do OCR nas pontas (símbolos, letras soltas, dígito solto no fim)."""
    toks = txt.split()
    while toks and not TOKEN_OK.match(toks[0]):
        toks.pop(0)
    while toks and (
        not TOKEN_OK.match(toks[-1])
        or (len(toks[-1]) == 1 and not toks[-1].isalpha())
        # letra solta depois de "LTDA"/"S.A." ("... S.A. É" do OCR)
        or (len(toks) > 1 and len(toks[-1]) == 1 and SUFIXO_RE.search(toks[-2]))
    ):
        toks.pop()
    return " ".join(toks).strip(" ,;:-")


def proxima_linha(linhas: list, i: int):
    """Próxima linha não vazia depois de i -> (índice, texto) ou (None, '')."""
    j = i + 1
    while j < len(linhas):
        if linhas[j]:
            return j, linhas[j]
        j += 1
    return None, ""


def cnj_valido(num: str) -> bool:
    """Confere os dígitos verificadores do nº CNJ (pega erro de OCR)."""
    m = re.match(r"(\d{7})-(\d{2})\.(\d{4})\.(\d)\.(\d{2})\.(\d{4})$", num)
    if not m:
        return False
    nnnnnnn, dd, aaaa, j, tr, oooo = m.groups()
    return 98 - (int(nnnnnnn + aaaa + j + tr + oooo + "00") % 97) == int(dd)


def nome_proprio(s: str) -> str:
    """JOSE DA SILVA -> Jose da Silva"""
    return re.sub(
        r"[A-Z0-9']+",
        lambda m: m.group(0).lower() if m.group(0) in PARTICULAS else m.group(0).capitalize(),
        s,
    )


def nome_hospital(destinatario: str) -> str:
    n = normalizar(destinatario)
    for trecho, curto in APELIDOS:
        if normalizar(trecho) in n:
            return curto
    return re.sub(r"\s+(LTDA|S/?A|S\.A\.?|EIRELI|EPP|ME)\W*$", "", n).strip()


# --------------------------------------------------------------------------
# Leitura do PDF: uma string por página (texto nativo ou OCR)
# --------------------------------------------------------------------------
_OCR = {"lang": "por"}


def ocr_imagem(img):
    import pytesseract

    try:
        return pytesseract.image_to_string(img, lang=_OCR["lang"], config="--psm 6")
    except pytesseract.TesseractError:
        if _OCR["lang"] == "eng":
            raise
        print("AVISO: idioma 'por' não encontrado no Tesseract; usando 'eng' (instale o pacote português).")
        _OCR["lang"] = "eng"
        return ocr_imagem(img)


def extrair_paginas(pdf, dpi=300, forcar_ocr=False, progresso=None) -> list:
    """pdf: caminho ou bytes. progresso(pagina, total) é chamado a cada página."""
    try:
        import pymupdf as fitz  # PyMuPDF (nome novo)
    except ImportError:
        import fitz  # PyMuPDF (versões antigas)
    from PIL import Image

    paginas = []
    doc = fitz.open(stream=pdf, filetype="pdf") if isinstance(pdf, (bytes, bytearray)) else fitz.open(pdf)
    with doc:
        for k, page in enumerate(doc, 1):
            if progresso:
                progresso(k, doc.page_count)
            txt = "" if forcar_ocr else page.get_text()
            if len(txt.strip()) < 50:  # página só com imagem -> OCR
                pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
                img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
                txt = ocr_imagem(img)
            paginas.append(txt)
    return paginas


# --------------------------------------------------------------------------
# Segmentação: página principal + envelope/continuação
# --------------------------------------------------------------------------
def eh_notificacao(texto: str) -> bool:
    n = sem_acentos(texto).upper()
    tem_reclamante = re.search(r"^\W*RECLAMANTE\s*:", n, re.M) is not None
    return tem_reclamante and (CNJ_RE.search(n) is not None or "PODER JUDICI" in n or "TRIBUNAL REGIONAL" in n)


def eh_envelope(texto: str) -> bool:
    n = sem_acentos(texto).upper()
    return any(s in n for s in SINAIS_ENVELOPE)


def segmentar(paginas: list) -> tuple:
    """-> (lista de notificações brutas, lista de páginas órfãs)."""
    registros, orfas, atual = [], [], None
    for num, texto in enumerate(paginas, 1):
        if eh_notificacao(texto):
            atual = {"paginas": [num], "texto": texto, "envelope": ""}
            registros.append(atual)
        elif atual is None:
            orfas.append(num)
        else:
            atual["paginas"].append(num)
            if eh_envelope(texto):
                atual["envelope"] += "\n" + texto
            else:
                atual["texto"] += "\n" + texto
    return registros, orfas


# --------------------------------------------------------------------------
# Extração dos campos de uma notificação
# --------------------------------------------------------------------------
def achar_processo(texto: str) -> str:
    achados = []
    for m in CNJ_RE.finditer(sem_acentos(texto)):
        dig = re.sub(r"\s", "", m.group(1))
        if len(dig) == 7:
            achados.append(f"{dig}-{m.group(2)}.{m.group(3)}.{m.group(4)}.{m.group(5)}.{m.group(6)}")
    for num in achados:
        if cnj_valido(num):
            return num
    return achados[0] if achados else ""


def achar_reclamante(linhas: list) -> str:
    for l in linhas:
        m = RECLAMANTE_RE.match(l)
        if m and m.group(1).strip():
            return limpar_ruido(m.group(1))
    return ""


def _fim_reclamado(s: str) -> bool:
    return bool(re.search(r"\(\s*\d+\s*\)$", s) or SUFIXO_RE.search(s))


def achar_reclamado_cabecalho(linhas: list):
    """-> (reclamado principal, nº de 'outros' ou None). O cabeçalho pode quebrar em várias linhas."""
    for i, l in enumerate(linhas):
        m = RECLAMADO_RE.match(l)
        if not m:
            continue
        s = limpar_ruido(m.group(1))
        j = i
        for _ in range(3):
            if s and _fim_reclamado(s):
                break
            j, prox = proxima_linha(linhas, j)
            if j is None:
                break
            if re.match(r"^\(?\d+\)?\W*$", prox) or (
                razao_maiuscula(prox) >= 0.7
                and sum(c.isalpha() for c in prox) >= 8
                and not re.match(r"^\W*(DESTINAT|NOTIFICA|AUDIENCIA|FICA)", prox)
            ):
                s = limpar_ruido(f"{s} {prox}")
            else:
                break
        if s and not OUTROS_RE.search(s):  # "E OUTROS (n)" pode ter quebrado de linha
            _, prox = proxima_linha(linhas, j if j is not None else i)
            if re.match(r"^\W*E\s+OUTROS", prox):
                s = limpar_ruido(f"{s} {prox}")
        mo = OUTROS_RE.search(s)
        outros = int(mo.group(1)) if mo else None
        principal = limpar_ruido(re.sub(r"\bE\s+OUTROS\b.*$", "", s)) if "OUTROS" in s else s
        return principal, outros
    return "", None


def achar_destinatario(linhas: list):
    """-> (nome, alertas). Corta endereço ' / RUA...' e 'na pessoa do responsável ...'."""
    for i, l in enumerate(linhas):
        m = DEST_RE.search(l)
        if not m:
            continue
        txt, alertas, cortou = m.group(1), [], False
        for padrao in (r"\s+/\s+", r"\s+NA PESSOA\b"):
            partes = re.split(padrao, txt, maxsplit=1)
            if len(partes) > 1:
                txt, cortou = partes[0], True
                if "PESSOA" in padrao:
                    alertas.append("Destinatário 'na pessoa do responsável' (conferir)")
        if not cortou and not SUFIXO_RE.search(txt):
            _, prox = proxima_linha(linhas, i)
            if (prox and not re.search(r"\d|CEP", prox)
                    and not re.match(r"^\W*(DESTINAT|NOTIFICA|AUDIENCIA|FICA)", prox)
                    and razao_maiuscula(prox) >= 0.7 and sum(c.isalpha() for c in prox) >= 8):
                txt = f"{txt} {prox}"
        nome = limpar_ruido(txt)
        if nome:
            return nome, alertas
    return "", []


def achar_destinatario_envelope(env_texto: str) -> str:
    """Plano B: nome impresso no envelope, logo antes da linha do endereço."""
    if not env_texto:
        return ""
    linhas = linhas_normalizadas(env_texto)
    ancora = next((i for i, l in enumerate(linhas) if "POSTAG" in l), None)
    inicio = (ancora + 1) if ancora is not None else 0
    for i in range(inicio, len(linhas)):
        if not ENDERECO_RE.match(linhas[i]):
            continue
        k = i - 1
        while k >= inicio and not linhas[k]:
            k -= 1
        if k < inicio:
            return ""
        nome = limpar_ruido(linhas[k])
        if SUFIXO_RE.search(nome) and len(nome.split()) <= 3:  # nome quebrado em 2 linhas
            k2 = k - 1
            while k2 >= inicio and not linhas[k2]:
                k2 -= 1
            if k2 >= inicio:
                prev = limpar_ruido(linhas[k2])
                if len(re.findall(r"[A-Z]{3,}", prev)) >= 2:
                    nome = f"{prev} {nome}"
        return nome
    return ""


def achar_audiencia(linhas: list):
    """Data/hora da audiência: 1ª data dd/mm/aaaa com hora (ignora data de postagem)."""
    uteis = "\n".join(l for l in linhas if "POSTAG" not in l)
    for m in DATA_HORA_RE.finditer(uteis):
        dia, mes, ano, hh, mm = map(int, m.groups())
        try:
            return datetime(ano, mes, dia), f"{hh:02d}:{mm:02d}"
        except ValueError:
            continue
    for m in DATA_RE.finditer(uteis):
        dia, mes, ano = map(int, m.groups())
        try:
            return datetime(ano, mes, dia), ""
        except ValueError:
            continue
    return None, ""


def achar_vara(flat: str, processo: str = "") -> str:
    """VARA. No TRT-1 (.5.01.) os 4 últimos dígitos do nº CNJ são a vara; mais confiável
    que o OCR do cabeçalho (que às vezes lê 'ª' como dígito). Texto é só plano B."""
    if processo and cnj_valido(processo) and processo.split(".")[2:4] == ["5", "01"]:
        return f"{int(processo[-4:])}ª VT"
    m = VARA_RE.search(flat)
    return f"{int(m.group(1))}ª VT" if m else ""


def _rotular_tipo(s: str):
    """-> (rótulo, era 'não una'). Em trecho com vários rótulos vale o mais à direita."""
    if re.search(r"\bNAO\s+UNA\b", s):
        return "Inicial", True
    padroes = (
        (r"\bINICIAL\b", "Inicial"),
        (r"\bINSTRU[CG]AO\b", "Instrução"),
        (r"\bCONCILIA[CG]AO\b", "Conciliação"),
        (r"\bUNA\b", "Una"),
    )
    achados = [(m.start(), rot) for pad, rot in padroes for m in re.finditer(pad, s)]
    return (max(achados)[1], False) if achados else ("", False)


def achar_tipo_audiencia(linhas: list, flat: str):
    """MODALIDADE: Una / Inicial / Instrução / Conciliação. -> (tipo, era 'não una')."""
    for l in linhas:  # 1) títulos curtos tipo "AUDIÊNCIA UNA - TELEPRESENCIAL"
        if "AUDIENCIA" in l and len(l) <= 80 and not DATA_RE.search(l):
            r = _rotular_tipo(l)
            if r[0]:
                return r
    m = DATA_HORA_RE.search(flat)  # 2) trecho logo antes da data/hora
    if m:
        r = _rotular_tipo(flat[max(0, m.start() - 120): m.start()])
        if r[0]:
            return r
    m = re.search(r"AUDIENCIA\s+(?:\w+\s+){0,2}?(NAO\s+UNA|UNA|INICIAL|INSTRU[CG]AO)\b", flat)
    return _rotular_tipo(m.group(1)) if m else ("", False)


def achar_formato(flat: str, tem_zoom: bool, tem_endereco: bool) -> str:
    """AUDIÊNCIAS: Presencial / Telepresencial / Híbrida."""
    if re.search(r"\bHIBRID[AO]\b", flat):
        return "Híbrida"
    if re.search(r"\bTELEPRESENCIAL\b|\bVIDEOCONFERENCIA\b", flat):
        return "Telepresencial"
    if re.search(r"\bPRESENCIAL\b", flat):
        return "Presencial"
    if tem_zoom:
        return "Telepresencial"
    return "Presencial" if tem_endereco else ""


def achar_testemunhas(flat: str) -> str:
    """Só preenche 'Não' quando o documento diz expressamente que não haverá testemunhas."""
    if re.search(
        r"NAO\s+SERA\s+PRODUZIDA\s+PROVA\s+TESTEMUNHAL|NAO\s+SERAO\s+OUVIDAS\s+TESTEMUNHAS"
        r"|NAO\s+HAVENDO\s+NECESSIDADE\s+DE\s+AS\s+PARTES\s+TRAZEREM", flat):
        return "Não"
    return ""


def achar_endereco_forum(flat: str) -> str:
    """Endereço da audiência (fórum, CEP 2023x), no formato da planilha."""
    for m in ENDERECO_FORUM_RE.finditer(flat):
        forum_conhecido = m.group(3) in ("LAVRADIO", "GOMES FREIRE")
        if not forum_conhecido and not re.search(r"CEP\W{0,3}2023\d", flat[m.end(): m.end() + 90]):
            continue
        tipo = "Avenida" if m.group(1) in ("AVENIDA", "AV") else "Rua"
        art = (m.group(2) or "").strip().lower()
        nome = " ".join(w.capitalize() for w in m.group(3).split())
        andar = f" - {m.group(5)}º andar" if m.group(5) else ""
        return f"{tipo} {art + ' ' if art else ''}{nome}, nº. {m.group(4)}{andar} - Centro - RJ"
    return ""


def achar_zoom(texto: str):
    """-> (link, 'ID da reunião: X – senha: Y')."""
    raw = re.sub(r"([-/?=])\s*\n\s*", r"\1", sem_acentos(texto))  # URL quebrada em 2 linhas
    m = ZOOM_RE.search(raw)
    link = m.group(0) if m else ""
    # O OCR costuma trocar o "1" de "trt1" por "l" ou "I" no subdomínio.
    link = re.sub(r"(?i)^(https?://)trt[1lI|]{1,2}-", r"\1trt1-", link)
    partes = []
    mi = ID_REUNIAO_RE.search(raw)
    if mi:
        d = re.sub(r"\D", "", mi.group(1))
        if len(d) in (9, 10, 11):  # formato do Zoom: 3-3-3, 3-3-4 ou 3-4-4
            meio = 4 if len(d) == 11 else 3
            partes.append(f"ID da reunião: {d[:3]} {d[3:3 + meio]} {d[3 + meio:]}")
    ms = SENHA_RE.search(raw)
    if ms:
        partes.append(f"senha: {ms.group(1)}")
    return link, " – ".join(partes)


def analisar(reg: dict) -> dict:
    linhas = linhas_normalizadas(reg["texto"])
    flat = " ".join(l for l in linhas if l)
    alertas = []

    processo = achar_processo(reg["texto"])
    reclamado, outros = achar_reclamado_cabecalho(linhas)
    dest, al = achar_destinatario(linhas)
    alertas += al
    fonte = "página"
    if not dest:
        dest = achar_destinatario_envelope(reg["envelope"])
        fonte = "envelope"
        if dest:
            alertas.append("Destinatário lido do envelope (conferir)")
    if not dest and reclamado:
        dest, fonte = reclamado, "cabeçalho"
        alertas.append("Destinatário não encontrado; usei o reclamado do cabeçalho (conferir)")

    data, hora = achar_audiencia(linhas)
    tipo, nao_una = achar_tipo_audiencia(linhas, flat)
    if nao_una:
        alertas.append("'Audiência não una' registrada como Inicial (conferir)")
    link, id_senha = achar_zoom(reg["texto"])
    endereco = achar_endereco_forum(flat)
    formato = achar_formato(flat, bool(link or id_senha), bool(endereco))

    partes_link = []
    if link or id_senha:
        partes_link.append("\n".join(p for p in (link, id_senha) if p))
        alertas.append("Link/ID/senha do Zoom lidos por OCR (conferir)")
    elif formato in ("Telepresencial", "Híbrida"):
        alertas.append("Link do Zoom não encontrado (preencher)")
    if endereco and formato in ("Presencial", "Híbrida"):
        partes_link.append(f"Comparecer: {endereco}")

    vara = achar_vara(flat, processo)
    faltando = [nome for nome, v in (("VARA", vara), ("MODALIDADE", tipo), ("AUDIÊNCIAS", formato)) if not v]
    if faltando:
        alertas.append("Não identificado: " + ", ".join(faltando))

    return {
        "processo": processo,
        "reclamante": achar_reclamante(linhas),
        "reclamado": reclamado,
        "outros": outros,
        "destinatario": dest,
        "fonte_dest": fonte if dest else "",
        "data": data,
        "hora": hora,
        "vara": vara,
        "tipo": tipo,
        "formato": formato,
        "testemunhas": achar_testemunhas(flat),
        "link_local": "\n".join(partes_link),
        "alertas": alertas,
        "paginas": reg["paginas"],
    }


# --------------------------------------------------------------------------
# Agrupamento: uma linha por processo
# --------------------------------------------------------------------------
def agrupar(notificacoes: list) -> list:
    grupos = {}
    for n in notificacoes:
        chave = n["processo"] or f"SEM-NUMERO::{n['arquivo']}::{n['paginas'][0]}"
        g = grupos.setdefault(chave, {
            "processo": n["processo"], "reclamante": "", "data": None, "hora": "",
            "polos": [], "reclamado": "", "outros": None, "vara": "", "tipo": "",
            "formato": "", "testemunhas": "", "link_local": "", "origem": [], "fontes": [],
            "alertas": [],
        })
        if g["reclamante"] and n["reclamante"] and normalizar(g["reclamante"]) != normalizar(n["reclamante"]):
            g["alertas"].append("Reclamantes diferentes no mesmo processo (OCR?)")
        g["reclamante"] = g["reclamante"] or n["reclamante"]
        if g["data"] and n["data"] and (g["data"], g["hora"]) != (n["data"], n["hora"]):
            g["alertas"].append("Datas/horas diferentes entre as notificações do processo")
        if not g["data"] and n["data"]:
            g["data"], g["hora"] = n["data"], n["hora"]
        if n["destinatario"]:
            conhecidos = {normalizar(p) for p in g["polos"]}
            if normalizar(n["destinatario"]) not in conhecidos:
                g["polos"].append(n["destinatario"])
        g["reclamado"] = g["reclamado"] or n["reclamado"]
        if n["outros"] is not None:
            g["outros"] = max(g["outros"] or 0, n["outros"])
        for campo in ("vara", "tipo", "formato", "testemunhas", "link_local"):
            g[campo] = g[campo] or n[campo]
        g["origem"].append(f"{n['arquivo']} p.{n['paginas'][0]}")
        g["fontes"].append({"arquivo": n["arquivo"], "pagina": n["paginas"][0]})
        g["alertas"] += n["alertas"]

    lista = list(grupos.values())
    for g in lista:
        if not g["processo"]:
            g["alertas"].append("Número do processo não encontrado")
        elif not cnj_valido(g["processo"]):
            g["alertas"].append("Número do processo com dígito verificador inválido (provável erro de OCR)")
        if not g["reclamante"]:
            g["alertas"].append("Reclamante não encontrado")
        if not g["data"]:
            g["alertas"].append("Data da audiência não encontrada")
        elif not g["hora"]:
            g["alertas"].append("Hora da audiência não encontrada")
        if not g["polos"]:
            g["alertas"].append("Hospital (destinatário) não encontrado")

    por_nome = {}  # mesmo reclamante em processos diferentes (só avisa, não une)
    for g in lista:
        if g["reclamante"]:
            por_nome.setdefault(normalizar(g["reclamante"]), []).append(g)
    for gs in por_nome.values():
        if len(gs) > 1:
            for g in gs:
                g["alertas"].append("Mesmo reclamante em mais de um processo (conferir)")
    try:  # nomes muito parecidos = possível erro de OCR (opcional)
        from rapidfuzz import fuzz

        nomes = list(por_nome)
        for a in range(len(nomes)):
            for b in range(a + 1, len(nomes)):
                if fuzz.ratio(nomes[a], nomes[b]) >= 90:
                    for g in por_nome[nomes[a]] + por_nome[nomes[b]]:
                        g["alertas"].append("Nome parecido com outro reclamante (conferir)")
    except ImportError:
        pass

    for g in lista:
        g["alertas"] = list(dict.fromkeys(g["alertas"]))
    lista.sort(key=lambda g: (g["data"] or datetime.max, g["hora"] or "99:99", normalizar(g["reclamante"])))
    return lista


# --------------------------------------------------------------------------
# Excel no formato do modelo "Audiências 2026"
# --------------------------------------------------------------------------
COLUNAS = ["STATUS", "DATA", "HORA", "HOSPITAL", "RECLAMANTE", "PROCESSO", "VARA",
           "MODALIDADE", "AUDIÊNCIAS", "TESTEMUNHAS", "OBSERVAÇÃO", "LINK E SENHA"]
LARGURAS = [12.1, 10.7, 8.1, 32.4, 48.3, 26.9, 20.3, 15.9, 16.1, 14.0, 27.7, 80.9]
FORMATO_HORA = r"[$-F400]h:mm:ss\ AM/PM"


def montar_linhas(grupos: list) -> list:
    """Grupos -> linhas da planilha (dicionários só com texto, fáceis de editar e
    de mandar como JSON para a página web). DATA = "AAAA-MM-DD", HORA = "HH:MM"."""
    linhas = []
    for g in grupos:
        linhas.append({
            "status": STATUS_PADRAO,
            "data": g["data"].strftime("%Y-%m-%d") if g["data"] else "",
            "hora": g["hora"],
            "hospital": "\n".join(dict.fromkeys(nome_hospital(p) for p in g["polos"])),
            "reclamante": nome_proprio(g["reclamante"]) if g["reclamante"] else "",
            "processo": g["processo"],
            "vara": g["vara"],
            "modalidade": g["tipo"],
            "audiencias": g["formato"],
            "testemunhas": g["testemunhas"],
            "observacao": "",
            "link": g["link_local"],
            # apoio (aba "Conferir")
            "reclamante_original": g["reclamante"],
            "destinatarios": g["polos"],
            "outros": g["outros"],
            "alertas": g["alertas"],
            "origem": g["origem"],
            "fontes": g.get("fontes", []),
        })
    return linhas


def _data(valor):
    if not valor:
        return None
    try:
        return datetime.strptime(str(valor)[:10], "%Y-%m-%d")
    except ValueError:
        return None


def _hora(valor):
    m = re.match(r"^(\d{1,2}):(\d{2})", str(valor or ""))
    if not m:
        return None
    try:
        return time(int(m.group(1)), int(m.group(2)))
    except ValueError:
        return None


def nome_arquivo_padrao(linhas: list) -> str:
    datas = [d for d in (_data(l.get("data")) for l in linhas) if d]
    if not datas:
        return "Audiências.xlsx"
    ini, fim = min(datas), max(datas)
    return f"Audiências_{ini.year}_-_{ini:%d_%m_%Y}_a_{fim:%d_%m_%Y}.xlsx"


def gerar_excel(linhas: list, saida):
    """Grava a planilha. 'saida' pode ser um caminho ou um arquivo aberto (BytesIO)."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.styles.colors import Color

    fino = Side(style="thin")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    amarelo = PatternFill("solid", start_color="FFF2CC")

    wb = Workbook()
    ws = wb.active
    anos = Counter(d.year for d in (_data(l.get("data")) for l in linhas) if d)
    ws.title = str(anos.most_common(1)[0][0]) if anos else "Audiências"
    ws.sheet_view.zoomScale = 85

    # cabeçalho (azul do tema, texto branco, centralizado)
    for c, nome in enumerate(COLUNAS, 1):
        cel = ws.cell(1, c, nome)
        cel.font = Font(name="Calibri", size=11, bold=True, color=Color(theme=0))
        cel.fill = PatternFill("solid", fgColor=Color(theme=4))
        cel.border = borda
        cel.alignment = Alignment(horizontal="left" if nome == "HOSPITAL" else "center",
                                  vertical="center", wrap_text=True if nome == "LINK E SENHA" else None)
    for c, larg in enumerate(LARGURAS, 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = larg

    centro = Alignment(horizontal="center", vertical="center")
    for r, l in enumerate(linhas, 2):
        valores = [
            l.get("status"), _data(l.get("data")), _hora(l.get("hora")), l.get("hospital"),
            l.get("reclamante"), l.get("processo"), l.get("vara"), l.get("modalidade"),
            l.get("audiencias"), l.get("testemunhas"), l.get("observacao"), l.get("link"),
        ]
        for c, v in enumerate(valores, 1):
            cel = ws.cell(r, c, v if v not in ("", None) else None)
            cel.border = borda
            cel.font = Font(name="Calibri", size=11, bold=(c == 1), color=Color(theme=1))
            if c == 4:
                cel.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif c in (5, 11):
                cel.alignment = Alignment(vertical="center")
            elif c == 12:
                cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cel.alignment = centro
        ws.cell(r, 2).number_format = "mm-dd-yy"
        ws.cell(r, 3).number_format = FORMATO_HORA
        alertas = l.get("alertas") or []
        if alertas:  # marca o nº do processo e explica no comentário
            ws.cell(r, 6).fill = amarelo
            ws.cell(r, 6).comment = Comment("Conferir:\n- " + "\n- ".join(alertas), "Leitor")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{max(len(linhas) + 1, 2)}"

    # aba de apoio: tudo que não cabe no modelo
    ws2 = wb.create_sheet("Conferir")
    cab2 = ["PROCESSO", "RECLAMANTE", "DESTINATÁRIOS (nome completo no documento)",
            "OUTROS RECLAMADOS (nº)", "ALERTAS", "ORIGEM"]
    ws2.append(cab2)
    for c in ws2[1]:
        c.font = Font(bold=True)
    for l in linhas:
        alertas = l.get("alertas") or []
        ws2.append([l.get("processo"), l.get("reclamante_original") or l.get("reclamante"),
                    "\n".join(l.get("destinatarios") or []), l.get("outros"),
                    "\n".join(alertas), "; ".join(l.get("origem") or [])])
        for c in (3, 5, 6):
            ws2.cell(ws2.max_row, c).alignment = Alignment(wrap_text=True, vertical="top")
        if alertas:
            for c in range(1, len(cab2) + 1):
                ws2.cell(ws2.max_row, c).fill = amarelo
    for col, larg in zip("ABCDEF", [28, 34, 60, 14, 55, 30]):
        ws2.column_dimensions[col].width = larg
    ws2.freeze_panes = "A2"
    wb.save(saida)


# --------------------------------------------------------------------------
# Principal
# --------------------------------------------------------------------------
def processar_paginas(paginas: list, nome_arquivo: str):
    registros, orfas = segmentar(paginas)
    notificacoes = []
    for reg in registros:
        n = analisar(reg)
        n["arquivo"] = nome_arquivo
        notificacoes.append(n)
    return notificacoes, orfas


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pasta", help="Pasta com os PDFs")
    ap.add_argument("-o", "--saida", help="Nome do Excel (padrão: Audiências_AAAA_-_dd_mm_aaaa_a_dd_mm_aaaa.xlsx)")
    ap.add_argument("--salvar-texto", action="store_true", help="Grava o texto lido de cada PDF em texto_lido/")
    ap.add_argument("--forcar-ocr", action="store_true", help="Faz OCR mesmo se o PDF tiver texto")
    ap.add_argument("--tesseract-cmd", help="Caminho do tesseract.exe (Windows)")
    ap.add_argument("--dpi", type=int, default=300)
    args = ap.parse_args()

    if args.tesseract_cmd:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = args.tesseract_cmd

    pasta = Path(args.pasta)
    pdfs = sorted({p for p in list(pasta.glob("*.pdf")) + list(pasta.glob("*.PDF"))})
    if not pdfs:
        sys.exit(f"Nenhum PDF encontrado em {pasta}")
    dir_texto = Path("texto_lido")
    if args.salvar_texto:
        dir_texto.mkdir(exist_ok=True)

    todas = []
    for k, pdf in enumerate(pdfs, 1):
        print(f"[{k}/{len(pdfs)}] {pdf.name}")
        try:
            paginas = extrair_paginas(pdf, dpi=args.dpi, forcar_ocr=args.forcar_ocr)
        except Exception as e:
            print(f"   ERRO ao ler: {e}")
            continue
        if args.salvar_texto:
            corpo = "\n".join(f"\n=== PÁGINA {i} ===\n{t}" for i, t in enumerate(paginas, 1))
            (dir_texto / f"{pdf.stem}.txt").write_text(corpo, encoding="utf-8")
        notifs, orfas = processar_paginas(paginas, pdf.name)
        todas += notifs
        print(f"   {len(paginas)} páginas, {len(notifs)} notificações")
        if orfas:
            print(f"   AVISO: páginas {orfas} não reconhecidas (antes de qualquer notificação)")

    linhas = montar_linhas(agrupar(todas))
    saida = Path(args.saida) if args.saida else Path(nome_arquivo_padrao(linhas))
    gerar_excel(linhas, saida)
    revisar = sum(1 for l in linhas if l["alertas"])
    print(f"\nPronto: {len(todas)} notificações -> {len(linhas)} processo(s) em {saida} ({revisar} com alerta).")


if __name__ == "__main__":
    main()
