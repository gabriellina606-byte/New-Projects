"""Testes do leitor_processos.py.

    python -m unittest discover -s testes -v

Os testes de ponta a ponta geram PDFs escaneados fictícios e precisam do
Tesseract instalado (com português); sem ele, são pulados.
"""
import io
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "testes"))

import leitor_processos as lp  # noqa: E402


class TestRegras(unittest.TestCase):
    def test_cnj_com_dois_pontos_do_ocr(self):
        self.assertEqual(lp.achar_processo("PROCESSO: 0100877-49.2026.5:01.0012 |"), "0100877-49.2026.5.01.0012")

    def test_cnj_com_espaco_no_meio(self):
        self.assertEqual(lp.achar_processo("0100 877-49.2026.5.01.0012"), "0100877-49.2026.5.01.0012")

    def test_cnj_prefere_numero_valido(self):
        texto = "0100877-48.2026.5.01.0012 e 0100877-49.2026.5.01.0012"
        self.assertEqual(lp.achar_processo(texto), "0100877-49.2026.5.01.0012")

    def test_vara_pelo_numero_do_processo(self):
        self.assertEqual(lp.achar_vara("122 VARA DO TRABALHO", "0100877-49.2026.5.01.0012"), "12ª VT")

    def test_vara_pelo_texto_quando_nao_e_trt1(self):
        self.assertEqual(lp.achar_vara("3A VARA DO TRABALHO DE NITEROI", ""), "3ª VT")

    def test_letra_solta_depois_de_sa(self):
        self.assertEqual(lp.limpar_ruido("HOSPITAL ILHA DO GOVERNADOR S.A. E"), "HOSPITAL ILHA DO GOVERNADOR S.A.")

    def test_e_outros_continua_intacto(self):
        linhas = lp.linhas_normalizadas("RECLAMADO: CASA DE PORTUGAL- HOSPITAL GERAL LTDA E OUTROS (3)")
        self.assertEqual(lp.achar_reclamado_cabecalho(linhas), ("CASA DE PORTUGAL- HOSPITAL GERAL LTDA", 3))

    def test_zoom_corrige_trt1_e_formata_id(self):
        link, id_senha = lp.achar_zoom(
            "https://trtl-jus-br.zoom.us/j/81234567890?pwd=AbC1\nID da reunião: 812 3456 7890\nSenha de acesso: 445566"
        )
        self.assertEqual(link, "https://trt1-jus-br.zoom.us/j/81234567890?pwd=AbC1")
        self.assertEqual(id_senha, "ID da reunião: 812 3456 7890 – senha: 445566")

    def test_zoom_url_quebrada_em_duas_linhas(self):
        link, _ = lp.achar_zoom("https://trt1-jus-br.zoom.us/j/\n81234567890")
        self.assertEqual(link, "https://trt1-jus-br.zoom.us/j/81234567890")

    def test_data_hora_ignora_postagem(self):
        linhas = lp.linhas_normalizadas("DATA DE POSTAGEM: 11/03/2026 10:00\npara o dia 22/04/2026 às 09h00")
        data, hora = lp.achar_audiencia(linhas)
        self.assertEqual((data.strftime("%d/%m/%Y"), hora), ("22/04/2026", "09:00"))

    def test_nao_una_vira_inicial_com_alerta(self):
        self.assertEqual(lp._rotular_tipo("AUDIENCIA NAO UNA"), ("Inicial", True))


class TestAgrupamentoEExcel(unittest.TestCase):
    def notif(self, **kw):
        base = {
            "processo": "0100877-49.2026.5.01.0012", "reclamante": "JOSE DA SILVA", "reclamado": "",
            "outros": None, "destinatario": "HOSPITAL ILHA DO GOVERNADOR S.A.", "fonte_dest": "página",
            "data": lp.datetime(2026, 4, 22), "hora": "09:00", "vara": "12ª VT", "tipo": "Inicial",
            "formato": "Presencial", "testemunhas": "", "link_local": "", "alertas": [], "paginas": [1],
            "arquivo": "a.pdf",
        }
        base.update(kw)
        return base

    def test_mesmo_processo_vira_uma_linha_com_hospitais_juntos(self):
        grupos = lp.agrupar([
            self.notif(),
            self.notif(destinatario="PRONTOCOR SERVICOS MEDICOS LTDA", arquivo="b.pdf", paginas=[3]),
        ])
        linhas = lp.montar_linhas(grupos)
        self.assertEqual(len(linhas), 1)
        self.assertEqual(linhas[0]["hospital"], "HOSPITAL ILHA DO GOVERNADOR\nPRONTO SOCORRO CLINICO PRONTOCOR")
        self.assertEqual(linhas[0]["fontes"], [{"arquivo": "a.pdf", "pagina": 1}, {"arquivo": "b.pdf", "pagina": 3}])

    def test_mesmo_reclamante_processos_diferentes_avisa(self):
        grupos = lp.agrupar([self.notif(), self.notif(processo="0100123-42.2025.5.01.0045")])
        self.assertEqual(len(grupos), 2)
        for g in grupos:
            self.assertIn("Mesmo reclamante em mais de um processo (conferir)", g["alertas"])

    def test_excel_em_memoria_e_nome_padrao(self):
        from openpyxl import load_workbook

        linhas = lp.montar_linhas(lp.agrupar([self.notif()]))
        linhas[0]["observacao"] = "Editado na tela"
        buf = io.BytesIO()
        lp.gerar_excel(linhas, buf)
        ws = load_workbook(io.BytesIO(buf.getvalue())).active
        self.assertEqual(ws.title, "2026")
        self.assertEqual([c.value for c in ws[1]], lp.COLUNAS)
        self.assertEqual(ws["B2"].value, lp.datetime(2026, 4, 22))
        self.assertEqual(ws["C2"].value, lp.time(9, 0))
        self.assertEqual(ws["K2"].value, "Editado na tela")
        self.assertEqual(lp.nome_arquivo_padrao(linhas), "Audiências_2026_-_22_04_2026_a_22_04_2026.xlsx")


@unittest.skipUnless(shutil.which("tesseract"), "Tesseract não instalado")
class TestPontaAPonta(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import gerar_exemplos

        cls.pasta = Path(tempfile.mkdtemp())
        gerar_exemplos.main(cls.pasta)
        todas = []
        for pdf in sorted(cls.pasta.glob("*.pdf")):
            notifs, orfas = lp.processar_paginas(lp.extrair_paginas(pdf), pdf.name)
            assert not orfas
            todas += notifs
        cls.linhas = {l["processo"]: l for l in lp.montar_linhas(lp.agrupar(todas))}

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.pasta, ignore_errors=True)

    def test_dois_processos(self):
        self.assertEqual(set(self.linhas), {"0100123-42.2025.5.01.0045", "0100877-49.2026.5.01.0012"})

    def test_telepresencial_com_zoom(self):
        l = self.linhas["0100123-42.2025.5.01.0045"]
        self.assertEqual((l["data"], l["hora"], l["vara"], l["modalidade"], l["audiencias"]),
                         ("2026-04-15", "10:30", "45ª VT", "Una", "Telepresencial"))
        self.assertEqual(l["hospital"], "HOSPITAL CASA DE PORTUGAL\nPRONTO SOCORRO CLINICO PRONTOCOR")
        self.assertIn("trt1-jus-br.zoom.us/j/81234567890", l["link"])
        self.assertEqual(l["outros"], 3)

    def test_presencial_no_forum(self):
        l = self.linhas["0100877-49.2026.5.01.0012"]
        self.assertEqual((l["data"], l["hora"], l["vara"], l["modalidade"], l["audiencias"], l["testemunhas"]),
                         ("2026-04-22", "09:00", "12ª VT", "Inicial", "Presencial", "Não"))
        self.assertEqual(l["reclamante"], "Jose Carlos da Silva Pereira")
        self.assertEqual(l["link"], "Comparecer: Rua do Lavradio, nº. 132 - 5º andar - Centro - RJ")
        self.assertEqual(l["alertas"], [])


if __name__ == "__main__":
    unittest.main()
